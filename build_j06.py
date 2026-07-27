#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
build_j06.py — Pipeline de atualizacao do acompanhamento de itens J06 (SBR4).

O QUE FAZ
---------
Le o arquivo-fonte GTO (List of Items) que fica na pasta ./data/ e regenera:
  * J06_Items_Analysis.xlsx  (planilha com cabecalho congelado + filtros)
  * J06_Dashboard.html       (dashboard responsivo)
Ambos passam a ter duas colunas de status:
  * Status Anterior  -> classificacao vigente ANTES da ultima atualizacao
  * Status Atual     -> classificacao vigente na fonte GTO mais recente

COMO ATUALIZAR (fluxo do usuario)
---------------------------------
1. Substitua ./data/GTO__LIST_OF_ITEMS.xlsx pelo arquivo novo (mesmo nome).
2. Rode:  python3 build_j06.py
O script guarda um historico de snapshots em ./data/status_history.json.
A cada execucao com dados diferentes, o snapshot atual vira "anterior"
automaticamente, sem intervencao manual.

FONTES DE DADOS (pasta ./data/)
-------------------------------
  GTO__LIST_OF_ITEMS.xlsx  Fonte primaria (status atual, descricao, certificado...).
  status_history.json      Historico de snapshots de status (base do Status Anterior).
  enrichment.json          Analise EN/PT, responsavel, categoria (curados a mao).
"""

import json
import re
import shutil
from datetime import date, datetime
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
SRC_XLSX = DATA / "GTO__LIST_OF_ITEMS.xlsx"
HISTORY_JSON = DATA / "status_history.json"
ENRICH_JSON = DATA / "enrichment.json"
TEMPLATE_HTML = ROOT / "templates" / "dashboard_template.html"
OUT_XLSX = ROOT / "J06_Items_Analysis.xlsx"
OUT_HTML = ROOT / "J06_Dashboard.html"

# Status considerados "encerrados" (nao entram nas pendencias abertas).
CLOSED_SET = {"1 - Validated by ICN", "2 - Not Blocking"}

# Cor de preenchimento por status (Excel) — herdado do arquivo original + novos.
STATUS_FILL = {
    "1 - Validated by ICN": "C6EFCE",
    "2 - Not Blocking": "E2EFDA",
    "3 - Blocking": "FFC7CE",
    "4 - Under Analysis": "FCE4D6",
    "4 - Under Analysis To Not Blocking": "FCE4D6",
    "5 - Waiting Proof": "FFEB9C",
    "6 - Waiting B05": "DDEBF7",
    "7 - Missing Vacuum Test or Sign": "FFF2CC",
    "8 - Awaiting full B05 completion": "D9D2E9",
}

CLOSED_MARKERS = ("Status closed", "Status encerrado", "no pending action")


# --------------------------------------------------------------------------- #
# Leitura da fonte GTO
# --------------------------------------------------------------------------- #
def read_gto(path: Path):
    """Le a planilha GTO. Retorna (registros_por_item, meta)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb.active
    header = {ws.cell(2, c).value: c for c in range(1, ws.max_column + 1)}

    def col(name):
        if name not in header:
            raise KeyError(f"Coluna '{name}' nao encontrada na fonte GTO.")
        return header[name]

    c_item, c_upd = col("Item"), col("Updated Status")
    c_ojx, c_ajx = col("Original Jx"), col("Actual Jx")
    c_desc, c_cert, c_ctx = col("Description"), col("Certificate"), col("Context")

    items = {}
    for r in range(3, ws.max_row + 1):
        it = ws.cell(r, c_item).value
        if it is None:
            continue
        items[str(it)] = {
            "item": it,
            "context": ws.cell(r, c_ctx).value or "",
            "ojx": ws.cell(r, c_ojx).value or "",
            "ajx": ws.cell(r, c_ajx).value or "",
            "desc": ws.cell(r, c_desc).value or "",
            "cert": ws.cell(r, c_cert).value or "",
            "status": (ws.cell(r, c_upd).value or "").strip(),
        }

    # metadados: contrato e "Gerado em" ficam na linha 1 (cantos)
    meta = {"contract": "", "generated": ""}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(1, c).value
        if isinstance(v, str):
            if v.startswith("Contrato"):
                meta["contract"] = v.split(":", 1)[-1].strip()
            elif v.startswith("Gerado"):
                meta["generated"] = v.replace("Gerado em", "").replace("às", "").strip()
    return items, meta


def norm_generated(raw: str) -> str:
    """'27/7/2026  10:42' -> '27/07/2026 10:42' (best-effort)."""
    if not raw:
        return date.today().strftime("%d/%m/%Y")
    m = re.search(r"(\d{1,2})/(\d{1,2})/(\d{4})(?:\D+(\d{1,2}):(\d{2}))?", raw)
    if not m:
        return raw
    d, mo, y = int(m.group(1)), int(m.group(2)), m.group(3)
    out = f"{d:02d}/{mo:02d}/{y}"
    if m.group(4):
        out += f" {int(m.group(4)):02d}:{m.group(5)}"
    return out


# --------------------------------------------------------------------------- #
# Historico de status (base do Status Anterior)
# --------------------------------------------------------------------------- #
def update_history(current_statuses: dict, source_label: str):
    """
    Acrescenta o snapshot atual ao historico se houver mudanca real.
    Retorna o dict {item: status} do snapshot ANTERIOR (para Status Anterior).
    """
    history = json.loads(HISTORY_JSON.read_text(encoding="utf-8"))
    snaps = history["snapshots"]
    last = snaps[-1]["statuses"]

    if current_statuses != last:
        snaps.append({
            "label": source_label,
            "date": date.today().isoformat(),
            "source": SRC_XLSX.name,
            "statuses": current_statuses,
        })
        previous = last
    else:
        # Reexecucao sem mudanca de dados: mantem historico, usa o penultimo.
        previous = snaps[-2]["statuses"] if len(snaps) >= 2 else last

    HISTORY_JSON.write_text(
        json.dumps(history, ensure_ascii=False, indent=1), encoding="utf-8"
    )
    return previous


# --------------------------------------------------------------------------- #
# Montagem dos registros finais
# --------------------------------------------------------------------------- #
def build_records(gto: dict, previous: dict, enrich: dict):
    records = []
    for key, g in gto.items():
        e = enrich.get(key, {})
        status = g["status"]
        closed = status in CLOSED_SET

        analysis = (e.get("analysis") or "").strip()
        analysis_pt = (e.get("analysis_pt") or "").strip()
        resp = (e.get("resp") or "").strip()
        category = e.get("category") or ("B05" if _looks_b05(g) else "General J06")
        note = e.get("note") or ""

        # Coerencia analise x status atual.
        if not closed and any(mk in analysis for mk in CLOSED_MARKERS):
            analysis, analysis_pt = "", ""
        if closed and not analysis:
            analysis = "Status closed - no pending action."
            analysis_pt = "Status encerrado - sem acao pendente."
        # Responsavel so faz sentido em item aberto.
        if not closed and (not resp or resp.startswith("—") or resp.startswith("-")):
            resp = ""
        if closed:
            resp = resp or "— (Closed – no pending)"

        records.append({
            "item": g["item"],
            "context": g["context"],
            "ojx": g["ojx"],
            "ajx": g["ajx"],
            "category": category,
            "desc": g["desc"],
            "cert": g["cert"],
            "status_prev": previous.get(key),   # Status Anterior
            "status": status,                    # Status Atual
            "analysis": analysis,
            "analysis_pt": analysis_pt,
            "resp": resp,
            "note": note,
            "closed": closed,
        })
    records.sort(key=lambda d: int(d["item"]))
    return records


def _looks_b05(g):
    blob = f"{g.get('desc','')} {g.get('cert','')}".upper()
    return "B05" in blob


# --------------------------------------------------------------------------- #
# Saida Excel
# --------------------------------------------------------------------------- #
HEADERS = [
    "Item", "Context", "OriginalJx", "ActualJx", "Category", "Description",
    "Certificate / Evidence", "Status Anterior", "Status Atual",
    "Analysis (EN)", "Responsible", "Analysis (PT)",
]
COL_WIDTHS = [7, 10, 11, 10, 12, 44, 24, 20, 20, 40, 22, 34]
NCOL = len(HEADERS)


def write_excel(records, last_updated, source_label):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Merged Analysis"

    thin = Side(style="thin", color="D0D7E5")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    title_fill = PatternFill("solid", fgColor="1F3864")
    hdr_fill = PatternFill("solid", fgColor="2E5496")

    last_col = get_column_letter(NCOL)

    # Linha 1 - titulo
    ws.merge_cells(f"A1:{last_col}1")
    t = ws["A1"]
    t.value = "J06 Items: Analysis, Status & Responsibility"
    t.font = Font(name="Arial", size=14, bold=True, color="FFFFFF")
    t.fill = title_fill
    t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 24

    # Linha 2 - subtitulo (esq) + data de atualizacao (canto dir)
    ws.merge_cells("A2:G2")
    s = ws["A2"]
    n_open = sum(1 for r in records if not r["closed"])
    s.value = (f"SBR4 · {source_label} · {len(records)} itens (ActualJx = J06) · "
               f"{n_open} pendencias abertas")
    s.font = Font(name="Arial", size=10, italic=True, color="44546A")
    s.alignment = Alignment(horizontal="left", vertical="center")
    ws.merge_cells(f"H2:{last_col}2")
    u = ws["H2"]
    u.value = f"Última atualização: {last_updated}"
    u.font = Font(name="Arial", size=10, bold=True, color="2E5496")
    u.alignment = Alignment(horizontal="right", vertical="center")
    ws.row_dimensions[2].height = 18

    # Linha 3 - cabecalhos
    for j, h in enumerate(HEADERS, start=1):
        c = ws.cell(3, j, h)
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c.fill = hdr_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws.row_dimensions[3].height = 30

    # Dados
    center = Alignment(horizontal="center", vertical="top", wrap_text=True)
    left = Alignment(horizontal="left", vertical="top", wrap_text=True)
    for i, d in enumerate(records):
        r = 4 + i
        changed = d["status_prev"] not in (None, "") and d["status_prev"] != d["status"]
        row = [
            d["item"], d["context"], d["ojx"], d["ajx"], d["category"], d["desc"],
            d["cert"],
            d["status_prev"] if d["status_prev"] else "— (novo item)",
            d["status"], d["analysis"], d["resp"], d["analysis_pt"],
        ]
        for j, val in enumerate(row, start=1):
            c = ws.cell(r, j, val)
            c.font = Font(name="Arial", size=10)
            c.border = border
            c.alignment = center if j in (1, 2, 3, 4, 5, 8, 9) else left
        # Preenchimento por status nas colunas 8 (Anterior) e 9 (Atual)
        prev_fill = STATUS_FILL.get(d["status_prev"])
        if prev_fill:
            ws.cell(r, 8).fill = PatternFill("solid", fgColor=prev_fill)
        cur_fill = STATUS_FILL.get(d["status"])
        if cur_fill:
            ws.cell(r, 9).fill = PatternFill("solid", fgColor=cur_fill)
        if changed:
            ws.cell(r, 9).font = Font(name="Arial", size=10, bold=True)

    # Larguras, congelamento e filtro
    for j, w in enumerate(COL_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = "A4"
    ws.auto_filter.ref = f"A3:{last_col}{3 + len(records)}"

    _write_summary(wb, records, last_updated)
    _write_changelog(wb, records, last_updated)
    wb.save(OUT_XLSX)


def _write_summary(wb, records, last_updated):
    ws = wb.create_sheet("Summary by Responsible")
    fmt = lambda r: r if (r and r.strip() and not r.startswith("—")) else "(To be assigned)"
    groups = {}
    for d in records:
        if d["closed"]:
            continue
        g = fmt(d["resp"])
        groups.setdefault(g, []).append(d)

    ws.merge_cells("A1:E1")
    ws["A1"] = "Open Items by Responsible (excludes closed: Validated / Not Blocking)"
    ws["A1"].font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")

    for j, h in enumerate(["Responsible", "Total", "Blocking", "Waiting/Other", "Items"], 1):
        c = ws.cell(2, j, h)
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E5496")
        c.alignment = Alignment(horizontal="center", vertical="center")

    order = sorted(groups, key=lambda g: -len(groups[g]))
    for i, g in enumerate(order):
        rows = groups[g]
        blk = sum(1 for d in rows if d["status"] == "3 - Blocking")
        items = ", ".join(str(d["item"]) for d in sorted(rows, key=lambda d: int(d["item"])))
        for j, val in enumerate([g, len(rows), blk, len(rows) - blk, items], 1):
            c = ws.cell(3 + i, j, val)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(horizontal="left" if j in (1, 5) else "center",
                                    vertical="top", wrap_text=(j == 5))
    for col, w in zip("ABCDE", [26, 9, 11, 14, 60]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"


def _write_changelog(wb, records, last_updated):
    """Aba com os itens que mudaram de status na ultima atualizacao."""
    changed = [d for d in records
               if d["status_prev"] not in (None, "") and d["status_prev"] != d["status"]]
    ws = wb.create_sheet("Status Changes")
    ws.merge_cells("A1:D1")
    ws["A1"] = f"Mudanças de status na última atualização ({last_updated}) — {len(changed)} item(ns)"
    ws["A1"].font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor="1F3864")
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    for j, h in enumerate(["Item", "Status Anterior", "Status Atual", "Description"], 1):
        c = ws.cell(2, j, h)
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="2E5496")
        c.alignment = Alignment(horizontal="center", vertical="center")
    for i, d in enumerate(sorted(changed, key=lambda d: int(d["item"]))):
        for j, val in enumerate([d["item"], d["status_prev"], d["status"], d["desc"]], 1):
            c = ws.cell(3 + i, j, val)
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(horizontal="center" if j in (1, 2, 3) else "left",
                                    vertical="top", wrap_text=(j == 4))
        if STATUS_FILL.get(d["status_prev"]):
            ws.cell(3 + i, 2).fill = PatternFill("solid", fgColor=STATUS_FILL[d["status_prev"]])
        if STATUS_FILL.get(d["status"]):
            ws.cell(3 + i, 3).fill = PatternFill("solid", fgColor=STATUS_FILL[d["status"]])
    for col, w in zip("ABCD", [8, 30, 30, 70]):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A3"


# --------------------------------------------------------------------------- #
# Saida HTML
# --------------------------------------------------------------------------- #
def write_html(records, last_updated, source_label):
    tpl = TEMPLATE_HTML.read_text(encoding="utf-8")
    payload = [{
        "item": d["item"], "ojx": d["ojx"], "ajx": d["ajx"], "desc": d["desc"],
        "cert": d["cert"], "category": d["category"],
        "status": d["status"], "status_prev": d["status_prev"],
        "analysis": d["analysis"], "analysis_pt": d["analysis_pt"],
        "resp": d["resp"], "note": d["note"], "closed": d["closed"],
    } for d in records]
    html = (tpl
            .replace("__DATA_JSON__", json.dumps(payload, ensure_ascii=False))
            .replace("__LAST_UPDATED__", last_updated)
            .replace("__ITEM_COUNT__", str(len(records)))
            .replace("__SOURCE_LABEL__", source_label))
    OUT_HTML.write_text(html, encoding="utf-8")


# --------------------------------------------------------------------------- #
def main():
    gto, meta = read_gto(SRC_XLSX)
    generated = norm_generated(meta.get("generated", ""))
    last_updated = generated
    source_label = "GTO List of Items"
    if meta.get("generated"):
        source_label += f" · gerado em {generated}"

    current = {k: v["status"] for k, v in gto.items()}
    previous = update_history(current, source_label)
    enrich = json.loads(ENRICH_JSON.read_text(encoding="utf-8"))

    records = build_records(gto, previous, enrich)
    write_excel(records, last_updated, source_label)
    write_html(records, last_updated, source_label)

    changed = sum(1 for d in records
                  if d["status_prev"] not in (None, "") and d["status_prev"] != d["status"])
    print(f"OK · {len(records)} itens · {changed} mudanca(s) de status · "
          f"atualizacao {last_updated}")
    print(f"  -> {OUT_XLSX.name}")
    print(f"  -> {OUT_HTML.name}")


if __name__ == "__main__":
    main()
